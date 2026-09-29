"""Build docs/INFER-user-guide.pptx — operator slides with UI / plan figures."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "docs" / "slides" / "assets"
OUT = ROOT / "docs" / "INFER-user-guide.pptx"

NAVY = RGBColor(0x0F, 0x17, 0x2A)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
ACCENT = RGBColor(0x1D, 0x4E, 0xD8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _set_run(run, text: str, *, size: int, bold: bool = False, color=NAVY) -> None:
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _bar(slide, prs) -> None:
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.08))
    sh.fill.solid()
    sh.fill.fore_color.rgb = ACCENT
    sh.line.fill.background()


def _footer(slide, prs, page: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.15), Inches(11.5), Inches(0.28))
    p = box.text_frame.paragraphs[0]
    _set_run(p.add_run(), "INFER  ·  internship POC  ·  not for life-safety sign-off", size=11, color=MUTED)
    r = box.text_frame.paragraphs[0].add_run()
    # page number on the right via a second box
    num = slide.shapes.add_textbox(Inches(12.2), Inches(7.15), Inches(0.7), Inches(0.28))
    np = num.text_frame.paragraphs[0]
    np.alignment = PP_ALIGN.RIGHT
    _set_run(np.add_run(), page, size=11, color=MUTED)


def _title(slide, text: str, top: float = 0.28) -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(0.55))
    p = box.text_frame.paragraphs[0]
    _set_run(p.add_run(), text, size=28, bold=True, color=NAVY)


def _sub(slide, text: str, top: float = 0.82) -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(0.4))
    p = box.text_frame.paragraphs[0]
    _set_run(p.add_run(), text, size=16, color=SLATE)


def _bullets(slide, items: list[str], left: float, top: float, width: float, height: float, *, size: int = 16) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        _set_run(p.add_run(), item, size=size, color=SLATE)


def _pic(slide, name: str, left, top, width, height) -> None:
    path = ASSETS / name
    if not path.is_file():
        return
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))


def _caption(slide, text: str, left: float, top: float, width: float) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.32))
    p = box.text_frame.paragraphs[0]
    _set_run(p.add_run(), text, size=12, color=MUTED)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    n = 0

    def next_page() -> str:
        nonlocal n
        n += 1
        return str(n)

    # 1 title
    s = _blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), prs.slide_height)
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT
    accent.line.fill.background()
    t = s.shapes.add_textbox(Inches(0.7), Inches(2.1), Inches(12), Inches(1))
    _set_run(t.text_frame.paragraphs[0].add_run(), "INFER", size=48, bold=True, color=WHITE)
    st = s.shapes.add_textbox(Inches(0.7), Inches(3.15), Inches(12), Inches(0.6))
    _set_run(st.text_frame.paragraphs[0].add_run(), "Intelligent Navigation and Facility Environment Reasoning", size=20, color=RGBColor(0xCB, 0xD5, 0xE1))
    sub = s.shapes.add_textbox(Inches(0.7), Inches(4.0), Inches(11), Inches(1.2))
    _set_run(
        sub.text_frame.paragraphs[0].add_run(),
        "User guide  ·  open a building, read 3D / plan / graph, route, what-if, share",
        size=18,
        color=RGBColor(0x94, 0xA3, 0xB8),
    )
    ft = s.shapes.add_textbox(Inches(0.7), Inches(6.6), Inches(11), Inches(0.4))
    _set_run(ft.text_frame.paragraphs[0].add_run(), "Internship proof-of-concept  ·  internal use", size=14, color=RGBColor(0x64, 0x74, 0x8B))

    # 2 what it does
    s = _blank(prs)
    _bar(s, prs)
    _title(s, "What INFER does")
    _sub(s, "Turns a static IFC / IndoorGML file into something you can route through — without changing the BIM.")
    _bullets(
        s,
        [
            "Ingest IFC or IndoorGML on this PC (spaces, doors, stairs, furniture).",
            "Show 3D BIM, a 2D plan, and a room connectivity graph.",
            "Right-click two points (or nearest exit) for a walkable path around desks.",
            "Double-click a door or a room to close it for this session.",
            "Evacuation heat + worst bottlenecks. Share GLB / QR (USDZ on iPhone).",
        ],
        0.55,
        1.4,
        12,
        4.8,
        size=20,
    )
    _footer(s, prs, next_page())

    # 3 workspace
    s = _blank(prs)
    _bar(s, prs)
    _title(s, "The workspace: three panes")
    _pic(s, "ui-workspace.png", 0.5, 1.15, 8.4, 5.5)
    _caption(s, "Real UI (before a model is loaded). Top bar: INFER · model menu · theme.", 0.5, 6.7, 8.4)
    _bullets(
        s,
        [
            "3D — orbit or fly; storey clip is independent of 2D.",
            "2D — IFC drawing or Navmesh routing.",
            "Graph — rooms as a network (not where you pick paths).",
            "Control tray — find / remove / restore by name.",
            "Refresh clears memory — open the file again.",
        ],
        9.1,
        1.3,
        3.8,
        5.4,
        size=15,
    )
    _footer(s, prs, next_page())

    # 4 open model
    s = _blank(prs)
    _bar(s, prs)
    _title(s, "Open a model")
    _sub(s, "Top bar → model name → Open model…  ·  or the Open model button on an empty pane.")
    _bullets(
        s,
        [
            "Accepts .ifc, .ifczip, .gml, .indoorgml.",
            "Pipeline: upload → extract → footprints → Geometry-rules graph.",
            "A large building can take minutes (furniture meshes).",
            "Watch the dialog progress and the status line.",
            "Files stay under data/ on this PC — not the cloud.",
        ],
        0.55,
        1.45,
        12,
        4.6,
        size=20,
    )
    _footer(s, prs, next_page())

    # 5 floor overview
    s = _blank(prs)
    _bar(s, prs)
    _title(s, "2D Viewer — the floorplan")
    _sub(s, "Same footprints the app draws. Example: Trapelo existing-conditions, Floor 2.")
    _pic(s, "plan-floor2-overview.png", 0.45, 1.25, 12.4, 5.55)
    _footer(s, prs, next_page())

    # 6 how to read
    s = _blank(prs)
    _bar(s, prs)
    _title(s, "How to read the plan")
    _pic(s, "plan-cubicles-closeup.png", 0.4, 1.15, 8.0, 5.5)
    _bullets(
        s,
        [
            "Grey fill — rooms (walkable space).",
            "Cream — true furniture occupancy, not a convex box.",
            "Chair wells stay empty — you can walk there.",
            "Amber glyphs — IFC doors (swing/slide when known).",
            "Blue disc — 3D camera on this storey.",
            "Legend toggles layers; furniture off ≠ ignored by routing.",
        ],
        8.6,
        1.3,
        4.3,
        5.4,
        size=16,
    )
    _footer(s, prs, next_page())

    # 7 3D (see the building before picking pins)
    s = _blank(prs)
    _bar(s, prs)
    _title(s, "3D Viewer")
    _sub(s, "Independent storey clip from the 2D dropdown. IndoorGML uses extruded footprints, not BIM.")
    _bullets(
        s,
        [
            "Orbit — drag around the building.",
            "Fly — WASD, Space up, Shift down, Ctrl faster.",
            "IFC tab — real mesh. Navmesh tab — portals / walkable overlay.",
            "Camera disc on the 2D plan tracks this viewpoint.",
        ],
        0.55,
        1.45,
        12.2,
        5.0,
        size=20,
    )
    _footer(s, prs, next_page())

    # 8 routing
    s = _blank(prs)
    _bar(s, prs)
    _title(s, "Ask for a route (Navmesh tab)")
    _pic(s, "plan-spaces-only.png", 0.4, 1.2, 7.6, 5.2)
    _bullets(
        s,
        [
            "Switch 2D from IFC (drawing) to Navmesh.",
            "Route: right-click start, then end (other floor is OK).",
            "Nearest exit: right-click once.",
            "Long right-click clears the path.",
            "Path stays ~40 cm off walls and goes around desks.",
        ],
        8.2,
        1.35,
        4.7,
        5.3,
        size=16,
    )
    _footer(s, prs, next_page())

    # 9 what-if
    s = _blank(prs)
    _bar(s, prs)
    _title(s, "What-if: block a door or a room")
    _sub(s, "Session overlay only — the IFC file is never rewritten.")
    _bullets(
        s,
        [
            "Navmesh: left-click selects (after a short delay).",
            "Double-click a portal — slash on the door; routing treats it as closed.",
            "Double-click a space — pale fill + slash; the whole room is closed.",
            "Double-click again to unblock, or use “N blocked · clear”.",
            "This is not Remove: Control / Graph right-click takes an entity out of the live model until you restore it.",
        ],
        0.55,
        1.35,
        12.2,
        5.2,
        size=18,
    )
    _footer(s, prs, next_page())

    # 10 evac
    s = _blank(prs)
    _bar(s, prs)
    _title(s, "Evacuation load")
    _bullets(
        s,
        [
            "Every unblocked room walks to its nearest exit (stairs/lifts included).",
            "Heat on doors and stairs; larger rooms count more (area proxy, not a headcount file).",
            "Blocking a door or room changes the heat.",
            "Worst bottlenecks list — click a row to jump 2D storey and fly 3D there.",
            "A planning sketch, not a fire-engineer’s model.",
        ],
        0.55,
        1.35,
        12.2,
        5.2,
        size=20,
    )
    _footer(s, prs, next_page())

    # 11 graph
    s = _blank(prs)
    _bar(s, prs)
    _title(s, "Graph Viewer")
    _sub(s, "Rooms as nodes. Most door nodes are hidden. Everyday paths: right-click the floorplan, not here.")
    _bullets(
        s,
        [
            "Geometry rules (default) — IFC plus healed doors / openings / stairs. Use this to route.",
            "IFC relations — authored links only; often too sparse on real exports.",
            "No TopologicPy in the UI — that API variant is an unfinished stub.",
            "Colours: yellow door heal, green opening, purple stair; grey = IFC. Red ring = nested parent.",
            "Right-click node/link to remove or disable; restore from Control.",
        ],
        0.55,
        1.35,
        12.2,
        5.2,
        size=18,
    )
    _footer(s, prs, next_page())

    # 12 control + share
    s = _blank(prs)
    _bar(s, prs)
    _title(s, "Control tray and sharing")
    _bullets(
        s,
        [
            "Control (right): filter by level / category, search, Remove / restore.",
            "When a navmesh route exists, Share on the top bar.",
            "GLB download; USDZ on iPhone; short-lived link + QR on the same network.",
            "Best 3D export if the 3D pane is open and IFC has finished loading.",
        ],
        0.55,
        1.4,
        12.2,
        5.0,
        size=20,
    )
    _footer(s, prs, next_page())

    # 13 limits
    s = _blank(prs)
    _bar(s, prs)
    _title(s, "Say this in a demo")
    _bullets(
        s,
        [
            "This is a POC. Do not treat paths or exits as certified life-safety.",
            "IFC quality matters — missing rooms or doors leave gaps.",
            "Refresh loses the UI session; derived files remain on disk.",
            "2D shows one storey; a multi-storey route draws the piece on the floor you are viewing.",
            "If it will not route: IFC tab, missed region, blocked door/room, or IFC-only graph.",
        ],
        0.55,
        1.4,
        12.2,
        5.2,
        size=20,
    )
    _footer(s, prs, next_page())

    # 14 recipe
    s = _blank(prs)
    _bar(s, prs)
    _title(s, "A five-minute walkthrough")
    _bullets(
        s,
        [
            "1. Open model… and wait until the toast shows spaces and links.",
            "2. 2D IFC tab → pick a storey → Legend if needed. Orbit or Fly in 3D.",
            "3. Navmesh → Route → right-click two rooms (try another floor for the second click).",
            "4. Nearest exit. Double-click a door, then a room; repeat the route.",
            "5. Evacuation load + bottlenecks. Share if you want the 3D hand-off.",
        ],
        0.55,
        1.4,
        12.2,
        5.2,
        size=20,
    )
    _footer(s, prs, next_page())

    prs.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
