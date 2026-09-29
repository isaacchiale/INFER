"""Build docs/INFER-user-guide.docx from docs/INFER-user-guide.md."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "docs" / "INFER-user-guide.md"
DOC_PATH = ROOT / "docs" / "INFER-user-guide.docx"


def _set_run_font(run, *, name: str = "Calibri", size: int = 11, bold: bool = False) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def _add_mixed(paragraph, text: str, *, size: int = 11) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            _set_run_font(run, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            _set_run_font(run, name="Consolas", size=size - 1)
        else:
            run = paragraph.add_run(part)
            _set_run_font(run, size=size)


def main() -> None:
    lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    i = 0
    in_code = False
    code_buf: list[str] = []
    table_buf: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_buf
        if not table_buf:
            return
        rows = table_buf
        table_buf = []
        if len(rows) < 2:
            return
        header, *body = rows
        # skip markdown separator row
        if body and all(re.fullmatch(r":?-+:?", c.strip()) for c in body[0]):
            body = body[1:]
        table = doc.add_table(rows=1 + len(body), cols=len(header))
        table.style = "Table Grid"
        for c, text in enumerate(header):
            table.rows[0].cells[c].text = text
        for r, row in enumerate(body, start=1):
            for c, text in enumerate(row):
                if c < len(table.rows[r].cells):
                    table.rows[r].cells[c].text = text
        doc.add_paragraph()

    def flush_code() -> None:
        nonlocal code_buf
        if not code_buf:
            return
        p = doc.add_paragraph()
        run = p.add_run("\n".join(code_buf))
        _set_run_font(run, name="Consolas", size=9)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(10)
        code_buf = []

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if line.startswith("```"):
            if in_code:
                flush_code()
                in_code = False
            else:
                flush_table()
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(raw)
            i += 1
            continue

        if line.startswith("|") and "|" in line[1:]:
            cells = [c.strip() for c in line.strip("|").split("|")]
            table_buf.append(cells)
            i += 1
            continue

        flush_table()

        if not line:
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("---"):
            pass
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_mixed(p, line[2:])
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            _add_mixed(p, re.sub(r"^\d+\.\s", "", line))
        elif line.startswith("*") and line.endswith("*") and not line.startswith("**"):
            p = doc.add_paragraph()
            run = p.add_run(line.strip("*"))
            _set_run_font(run, size=10)
            run.italic = True
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        else:
            p = doc.add_paragraph()
            _add_mixed(p, line)

        i += 1

    flush_table()
    flush_code()
    doc.save(DOC_PATH)
    print(f"Wrote {DOC_PATH}")


if __name__ == "__main__":
    main()
